"""
Приклад тестування функціоналу імпорту даних для oiHromadaSurvey

Цей скрипт демонструє як:
1. Створити тестові дані
2. Експортувати шаблон
3. Імітувати редагування
4. Імпортувати дані назад
5. Перевірити результати

Використання:
    bench execute orange_inventory.test_import_example.test_import_workflow
"""

import frappe
from frappe.utils import random_string


def create_test_organization():
	"""Створює тестову організацію"""
	org_name = f"TEST-ORG-{random_string(5)}"

	if frappe.db.exists("oiOrganization", org_name):
		return org_name

	org = frappe.get_doc(
		{
			"doctype": "oiOrganization",
			"name": org_name,
			"organization_name": f"Тестова організація {random_string(3)}",
			"abbreviation": f"ТЕСТ-{random_string(3)}",
			"tax_code": f"{frappe.utils.cint(frappe.utils.random_string(8))}",
			"enabled": 1,
		}
	)
	org.insert(ignore_permissions=True)
	frappe.db.commit()

	print(f"✅ Створено тестову організацію: {org_name}")
	return org_name


def create_test_survey_data(org_name):
	"""Створює тестові показники для організації"""

	# Створюємо групу
	group_id = f"test-group-{random_string(5)}"
	group = frappe.get_doc(
		{
			"doctype": "oiHromadaSurvey",
			"id": group_id,
			"title": "Тестова група показників",
			"type": "Група",
			"is_group": 1,
			"enabled": 1,
		}
	)
	group.insert(ignore_permissions=True)

	# Створюємо кількісний показник
	quant_id = f"test-quant-{random_string(5)}"
	quant = frappe.get_doc(
		{
			"doctype": "oiHromadaSurvey",
			"id": quant_id,
			"title": "Кількість комп'ютерів (тест)",
			"type": "Кількісні дані",
			"int_data": 10,
			"frequency": "Раз на квартал",
			"master_info": org_name,
			"parent_oihromadasurvey": group_id,
			"enabled": 1,
			"only_admin": 0,
		}
	)
	quant.insert(ignore_permissions=True)

	# Створюємо якісний показник
	qual_id = f"test-qual-{random_string(5)}"
	qual = frappe.get_doc(
		{
			"doctype": "oiHromadaSurvey",
			"id": qual_id,
			"title": "Наявність інтернету (тест)",
			"type": "Якісні дані",
			"bool_data": 0,
			"frequency": "Раз на рік",
			"master_info": org_name,
			"parent_oihromadasurvey": group_id,
			"enabled": 1,
			"only_admin": 0,
		}
	)
	qual.insert(ignore_permissions=True)

	frappe.db.commit()

	print("✅ Створено тестові показники:")
	print(f"   - Група: {group_id}")
	print(f"   - Кількісний: {quant_id} (поточне значення: 10)")
	print(f"   - Якісний: {qual_id} (поточне значення: Ні)")

	return {"group": group_id, "quant": quant_id, "qual": qual_id}


def export_template(org_name):
	"""Експортує шаблон для організації"""
	from orange_inventory.orange_inventory.doctype.oihromadasurvey.oihromadasurvey import export_org_template

	result = export_org_template(org=org_name, period="2025-Q4")
	file_url = result.get("file_url")

	print(f"✅ Експортовано шаблон: {file_url}")
	return file_url


def simulate_excel_edit_and_import(file_url, org_name, test_ids):
	"""
	Імітує редагування Excel файлу та імпорт.
	В реальності користувач відкриває файл, редагує колонку G, і завантажує назад.
	Тут ми імітуємо це програмно.
	"""
	import os

	from openpyxl import load_workbook

	# Отримуємо шлях до файлу
	file_doc = frappe.get_doc("File", {"file_url": file_url})
	file_path = file_doc.get_full_path()

	print(f"📝 Відкриваємо файл: {file_path}")

	# Відкриваємо та редагуємо
	wb = load_workbook(file_path)
	ws = wb["Показники"]

	# Знаходимо рядки з нашими тестовими ID та змінюємо значення
	changes_made = []
	for row_idx, row in enumerate(ws.iter_rows(min_row=5, values_only=False), start=5):
		survey_id = row[1].value  # Колонка B - ID

		if survey_id == test_ids["quant"]:
			# Змінюємо кількісне значення з 10 на 25
			old_val = row[6].value
			row[6].value = 25
			changes_made.append(f"   {survey_id}: {old_val} → 25")
			print(f"✏️  Змінено рядок {row_idx}: {survey_id} = 25")

		elif survey_id == test_ids["qual"]:
			# Змінюємо якісне значення з "Ні" на "Так"
			old_val = row[6].value
			row[6].value = "Так"
			changes_made.append(f"   {survey_id}: {old_val} → Так")
			print(f"✏️  Змінено рядок {row_idx}: {survey_id} = Так")

	# Зберігаємо зміни
	wb.save(file_path)
	print("💾 Збережено зміни в файл")

	# Тепер імпортуємо
	from orange_inventory.orange_inventory.doctype.oihromadasurvey.oihromadasurvey import import_from_excel

	print("\n📥 Початок імпорту...")
	result = import_from_excel(file_url=file_url, org=org_name)

	print("\n✅ Імпорт завершено!")
	print(f"   Статус: {result['status']}")
	print(f"   Оновлено: {result['updated']}")
	print(f"   Пропущено: {result['skipped']}")
	print(f"   Помилок: {len(result['errors'])}")

	if result["errors"]:
		print("\n❌ Помилки:")
		for error in result["errors"]:
			print(f"   - {error}")

	if result["details"]:
		print("\n📊 Деталі змін:")
		for detail in result["details"]:
			print(f"   - {detail['id']}: {detail['old_value']} → {detail['new_value']}")

	return result


def verify_changes(test_ids):
	"""Перевіряє що зміни були застосовані"""
	print("\n🔍 Перевірка результатів...")

	# Перевіряємо кількісний показник
	quant = frappe.get_doc("oiHromadaSurvey", test_ids["quant"])
	assert quant.int_data == 25, f"❌ Кількісний показник не оновився: {quant.int_data} != 25"
	print(f"   ✅ Кількісний показник: {quant.int_data} (очікувалось 25)")

	# Перевіряємо якісний показник
	qual = frappe.get_doc("oiHromadaSurvey", test_ids["qual"])
	assert qual.bool_data == 1, f"❌ Якісний показник не оновився: {qual.bool_data} != 1"
	print(f"   ✅ Якісний показник: {'Так' if qual.bool_data else 'Ні'} (очікувалось Так)")

	# Перевіряємо історію для кількісного
	quant_history = frappe.get_all(
		"oiHromadaSurveyHistory",
		filters={"parent": test_ids["quant"]},
		fields=["recorded_date", "period", "int_value"],
		order_by="recorded_date desc",
		limit=1,
	)

	if quant_history:
		print(f"   ✅ Історія створена для кількісного показника: {quant_history[0]}")
	else:
		print("   ⚠️  Історія не знайдена для кількісного показника")

	print("\n✅ Всі перевірки пройдено успішно!")


def cleanup_test_data(org_name, test_ids):
	"""Очищає тестові дані"""
	print("\n🧹 Очищення тестових даних...")

	try:
		# Видаляємо показники
		for doc_id in [test_ids["quant"], test_ids["qual"], test_ids["group"]]:
			if frappe.db.exists("oiHromadaSurvey", doc_id):
				frappe.delete_doc("oiHromadaSurvey", doc_id, force=1)
				print(f"   ✅ Видалено показник: {doc_id}")

		# Видаляємо організацію
		if frappe.db.exists("oiOrganization", org_name):
			frappe.delete_doc("oiOrganization", org_name, force=1)
			print(f"   ✅ Видалено організацію: {org_name}")

		frappe.db.commit()
		print("✅ Очищення завершено")

	except Exception as e:
		print(f"⚠️  Помилка при очищенні: {str(e)}")


def test_import_workflow():
	"""Основний тестовий workflow"""
	print("\n" + "=" * 60)
	print("🚀 ТЕСТУВАННЯ ФУНКЦІОНАЛУ ІМПОРТУ ДАНИХ")
	print("=" * 60 + "\n")

	org_name = None
	test_ids = None

	try:
		# 1. Створення тестової організації
		print("📋 Крок 1: Створення тестової організації")
		org_name = create_test_organization()

		# 2. Створення тестових даних
		print("\n📋 Крок 2: Створення тестових показників")
		test_ids = create_test_survey_data(org_name)

		# 3. Експорт шаблону
		print("\n📋 Крок 3: Експорт шаблону")
		file_url = export_template(org_name)

		# 4. Імітація редагування та імпорт
		print("\n📋 Крок 4: Редагування та імпорт")
		simulate_excel_edit_and_import(file_url, org_name, test_ids)

		# 5. Перевірка результатів
		print("\n📋 Крок 5: Перевірка результатів")
		verify_changes(test_ids)

		print("\n" + "=" * 60)
		print("✅ ТЕСТ ПРОЙДЕНО УСПІШНО!")
		print("=" * 60 + "\n")

		# 6. Очищення (опціонально, закоментуйте якщо хочете залишити дані)
		# cleanup_test_data(org_name, test_ids)

	except Exception as e:
		print("\n" + "=" * 60)
		print(f"❌ ТЕСТ ПРОВАЛЕНО: {str(e)}")
		print("=" * 60 + "\n")
		import traceback

		traceback.print_exc()

		# Очищення при помилці
		if org_name and test_ids:
			cleanup_test_data(org_name, test_ids)

		raise


if __name__ == "__main__":
	test_import_workflow()
